"""Text- und Eigenschafts-Extraktion je Dateityp.

Jeder Extraktor liefert `Extracted`: Textsegmente mit Fundstelle (Seite, Folie, Blatt)
plus Dokumenteigenschaften (Titel, Autor, Erstell-/Änderungsdatum).
Binärformate ohne lesbaren Text werden nur mit Metadaten erfasst.
"""
from __future__ import annotations

import csv
import io
import logging
import zipfile
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
import re

log = logging.getLogger(__name__)


@dataclass
class Segment:
    text: str
    locator: str | None = None
    page: int | None = None
    heading: str | None = None   # Abschnitt, zu dem der Text gehört (für DoCO/DEO)


@dataclass
class Extracted:
    segments: list[Segment] = field(default_factory=list)
    title: str | None = None
    authors: list[str] = field(default_factory=list)
    created: datetime | None = None
    modified: datetime | None = None
    pages: int | None = None
    note: str | None = None   # Hinweis, falls kein Text extrahiert wurde

    @property
    def text(self) -> str:
        return "\n\n".join(s.text for s in self.segments if s.text.strip())


MIME = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "doc": "application/msword",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "ppt": "application/vnd.ms-powerpoint",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "csv": "text/csv",
    "txt": "text/plain",
    "md": "text/markdown",
    "jpg": "image/jpeg",
    "jpeg": "image/jpeg",
    "png": "image/png",
    "sav": "application/x-spss-sav",
    "sps": "text/plain",
    "sas": "text/plain",
}

TEXT_EXT = {"txt", "md", "sps", "sas", "sts", "mdmt", "r", "py", "do", "tex", "bib"}


def _utc(dt) -> datetime | None:
    if not isinstance(dt, datetime):
        return None
    return dt if dt.tzinfo else dt.replace(tzinfo=timezone.utc)


def _pdf_date(raw) -> datetime | None:
    # Format: D:20180409144400+02'00'
    if not raw:
        return None
    s = str(raw).strip()
    if s.startswith("D:"):
        s = s[2:]
    for fmt, n in (("%Y%m%d%H%M%S", 14), ("%Y%m%d%H%M", 12), ("%Y%m%d", 8), ("%Y", 4)):
        try:
            return datetime.strptime(s[:n], fmt).replace(tzinfo=timezone.utc)
        except ValueError:
            continue
    return None


def _split_authors(raw: str | None) -> list[str]:
    if not raw:
        return []
    parts = [p.strip() for p in raw.replace(" and ", ";").replace(",", ";").split(";")]
    return [p for p in parts if p and p.lower() not in {"author", "user", "admin", "unknown"}]


def split_by_headings(text: str, current: str | None) -> list[tuple[str | None, str]]:
    """Teilt Seitentext an erkannten Abschnittsüberschriften; liefert (Überschrift, Text)-Paare.
    Der Abschnitt der Vorseite läuft weiter, bis eine neue Überschrift kommt."""
    from .deo import detect_heading

    out: list[tuple[str | None, list[str]]] = [(current, [])]
    for line in text.splitlines():
        h = detect_heading(line)
        if h:
            out.append((h, [line]))
        else:
            out[-1][1].append(line)
    return [(h, "\n".join(lines)) for h, lines in out if any(l.strip() for l in lines)]


# ------------------------------------------------------------------ Extraktoren
# REQ-49: personalisierte Lizenzstempel von Normen-Downloads (Name, Kundennummer, Download-ID, Zeitpunkt) gehören weder
# in den Index noch in Prompts an das Sprachmodell – z. B. „Normen-Download-DIN Media-<Name>-KdNr.<n>-ID.<id>-2026-10-29 23:53:33“.
_LICENCE_STAMP = re.compile(r"(?:Normen-Download|Norm-Download|Standards? download)-[^\n]{0,240}?(?:KdNr|Kd\.-?Nr|Customer No)\.?\s*\d+"
                            r"(?:-ID\.?\s*[\w]+)?(?:-\d{4}-\d{2}-\d{2}(?:\s+\d{2}:\d{2}(?::\d{2})?)?)?", re.I)


def strip_licence_stamps(text: str) -> str:
    return _LICENCE_STAMP.sub("", text) if text else text


def extract_pdf(path: Path) -> Extracted:
    from pypdf import PdfReader

    out = Extracted()
    reader = PdfReader(str(path))
    meta = reader.metadata or {}
    out.title = (meta.get("/Title") or "").strip() or None
    out.authors = _split_authors(meta.get("/Author"))
    out.created = _pdf_date(meta.get("/CreationDate"))
    out.modified = _pdf_date(meta.get("/ModDate"))
    out.pages = len(reader.pages)
    heading = None
    for i, page in enumerate(reader.pages, start=1):
        try:
            t = strip_licence_stamps(page.extract_text() or "")
        except Exception as e:  # defekte Seiten überspringen
            log.warning("PDF-Seite %s in %s: %s", i, path.name, e)
            t = ""
        if t.strip():
            for h, part in split_by_headings(t, heading):
                heading = h
                if part.strip():
                    out.segments.append(Segment(part, f"S. {i}", i, h))
    if not out.segments:
        out.note = "No extractable text (possibly a scan; OCR needed)"
    return out


def extract_docx(path: Path) -> Extracted:
    import docx

    out = Extracted()
    d = docx.Document(str(path))
    cp = d.core_properties
    out.title = (cp.title or "").strip() or None
    out.authors = _split_authors(cp.author)
    out.created, out.modified = _utc(cp.created), _utc(cp.modified)
    from .deo import clean_heading, detect_heading

    heading, buf = None, []

    def flush():
        if buf:
            out.segments.append(Segment("\n".join(buf), heading, None, heading))
            buf.clear()

    for p in d.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        style = (getattr(p.style, "name", "") or "").lower()
        is_head = style.startswith(("heading", "überschrift", "title", "titel")) and len(txt) <= 120
        h = clean_heading(txt) if is_head else detect_heading(txt)
        if h:
            flush()
            heading = h
        buf.append(txt)
    rows = []
    for t in d.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
    flush()
    if rows:
        out.segments.append(Segment("\n".join(rows), "Tables", None, heading))
    return out


def extract_pptx(path: Path) -> Extracted:
    from pptx import Presentation

    out = Extracted()
    prs = Presentation(str(path))
    cp = prs.core_properties
    out.title = (cp.title or "").strip() or None
    out.authors = _split_authors(cp.author)
    out.created, out.modified = _utc(cp.created), _utc(cp.modified)
    out.pages = len(prs.slides)
    for i, slide in enumerate(prs.slides, start=1):
        texts = []
        for shp in slide.shapes:
            if getattr(shp, "has_text_frame", False) and shp.text_frame.text.strip():
                texts.append(shp.text_frame.text.strip())
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            texts.append("Notizen: " + slide.notes_slide.notes_text_frame.text.strip())
        if texts:
            out.segments.append(Segment("\n".join(texts), f"Folie {i}", i))
    return out


def extract_xlsx(path: Path, max_rows: int = 200) -> Extracted:
    import openpyxl

    out = Extracted()
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    p = wb.properties
    out.title = (p.title or "").strip() or None
    out.authors = _split_authors(p.creator)
    out.created, out.modified = _utc(p.created), _utc(p.modified)
    for ws in wb.worksheets:
        rows = []
        for r_i, row in enumerate(ws.iter_rows(values_only=True)):
            if r_i >= max_rows:
                rows.append(f"… ({ws.max_row} Zeilen gesamt)")
                break
            vals = ["" if v is None else str(v) for v in row]
            if any(vals):
                rows.append(" | ".join(vals))
        if rows:
            out.segments.append(Segment(f"Blatt {ws.title}\n" + "\n".join(rows), f"Blatt {ws.title}"))
    wb.close()
    return out


def extract_csv(path: Path, max_rows: int = 200) -> Extracted:
    raw = path.read_bytes()
    text = raw.decode("utf-8", errors="ignore") if raw[:3] == b"\xef\xbb\xbf" else raw.decode("cp1252", errors="ignore")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t")
        rows = list(csv.reader(io.StringIO(text), dialect))
    except csv.Error:
        rows = [line.split(";") for line in text.splitlines()]
    body = "\n".join(" | ".join(r) for r in rows[:max_rows])
    header = f"Spalten: {', '.join(rows[0])}\nZeilen: {len(rows) - 1}\n" if rows else ""
    return Extracted(segments=[Segment(header + body)] if rows else [])


def extract_sav(path: Path) -> Extracted:
    """SPSS-Datensatz: Variablennamen + Labels sind für die Suche wertvoll, die Werte nicht."""
    try:
        import pyreadstat
    except ImportError:
        return Extracted(note="pyreadstat not installed – metadata only")
    try:
        _, meta = pyreadstat.read_sav(str(path), metadataonly=True)
    except Exception as e:
        return Extracted(note=f"SPSS file not readable: {e}")
    lines = [f"SPSS dataset, {meta.number_rows} cases, {meta.number_columns} variables"]
    for name in meta.column_names:
        label = meta.column_names_to_labels.get(name) or ""
        lines.append(f"{name}: {label}".rstrip(": "))
    return Extracted(segments=[Segment("\n".join(lines))], title=meta.file_label or None)


def extract_spv(path: Path) -> Extracted:
    """SPSS-Ausgabe (.spv) ist ein ZIP mit XML – Überschriften und Tabellentexte herausziehen."""
    import re

    try:
        with zipfile.ZipFile(path) as z:
            texts = []
            for n in z.namelist():
                if n.endswith(".xml") and ("heading" in n or "outputViewer" in n):
                    xml = z.read(n).decode("utf-8", errors="ignore")
                    texts += re.findall(r"<(?:vtx:)?(?:label|text)[^>]*>([^<]{3,})<", xml)
    except zipfile.BadZipFile:
        return Extracted(note="SPSS output not readable")
    uniq = list(dict.fromkeys(t.strip() for t in texts if t.strip()))
    return Extracted(segments=[Segment("\n".join(uniq[:2000]))] if uniq else [], note=None if uniq else "No text in SPSS output")


def extract_text(path: Path) -> Extracted:
    raw = path.read_bytes()
    if b"\x00" in raw[:2048]:
        return Extracted(note="Binary file – metadata only")
    for enc in ("utf-8", "cp1252", "latin-1"):
        try:
            return Extracted(segments=[Segment(raw.decode(enc))])
        except UnicodeDecodeError:
            continue
    return Extracted(note="Unknown encoding")


def extract_legacy_office(path: Path) -> Extracted:
    """.doc/.ppt: mit LibreOffice nach docx/pptx konvertieren, falls installiert."""
    import shutil
    import subprocess
    import tempfile

    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return Extracted(note="Legacy Office format – LibreOffice needed for text")
    target = {"doc": "docx", "ppt": "pptx", "xls": "xlsx"}[path.suffix.lower().lstrip(".")]
    with tempfile.TemporaryDirectory() as tmp:
        try:
            subprocess.run([soffice, "--headless", "--convert-to", target, "--outdir", tmp, str(path)],
                           check=True, capture_output=True, timeout=120)
        except Exception as e:
            return Extracted(note=f"Conversion failed: {e}")
        conv = Path(tmp) / (path.stem + "." + target)
        if not conv.exists():
            return Extracted(note="Conversion produced no output")
        res = EXTRACTORS[target](conv)
        # Datumsangaben stammen sonst von der Konvertierung, nicht vom Original
        res.created = res.modified = None
        return res


EXTRACTORS = {
    "pdf": extract_pdf,
    "docx": extract_docx,
    "pptx": extract_pptx,
    "xlsx": extract_xlsx,
    "xlsm": extract_xlsx,
    "csv": extract_csv,
    "sav": extract_sav,
    "spv": extract_spv,
    "doc": extract_legacy_office,
    "ppt": extract_legacy_office,
    "xls": extract_legacy_office,
}


def filetype_of(path: Path) -> str:
    ext = path.suffix.lower().lstrip(".")
    return ext or "unknown"


def extract(path: Path) -> Extracted:
    ft = filetype_of(path)
    fn = EXTRACTORS.get(ft) or (extract_text if ft in TEXT_EXT else None)
    if fn is None:
        return Extracted(note=f"File type .{ft} – metadata only")
    try:
        return fn(path)
    except Exception as e:
        log.exception("Extraktion fehlgeschlagen: %s", path)
        return Extracted(note=f"Extraction failed: {type(e).__name__}: {e}")
